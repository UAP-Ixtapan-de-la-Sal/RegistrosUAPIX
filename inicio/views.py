import os #Owen
from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required, user_passes_test
from django.utils import timezone #Owen
from django.conf import settings #Owen
from django.http import HttpResponse #Owen
from openpyxl import load_workbook #Owen
from .models import UsuarioSistema
from openpyxl.drawing.image import Image  # Para insertar imágenes en Excel
from .forms import RegistroUsuarioForm  # Importamos tu nuevo escudo de seguridad
from .models import (Registro) #Owen

# AQUI EMPEZO A MODIFICAR OWEN
def exportar_auditoria_excel(request):
    ruta_plantilla = os.path.join(settings.BASE_DIR, 'inicio', 'plantillas', 'plantilla_auditoria.xlsx')
    workbook = load_workbook(ruta_plantilla)
    hoja = workbook.active

    # === 1. TÍTULO DINÁMICO ===
    fecha_actual = timezone.localtime().strftime('%d/%m/%Y')
    
    # Celda del título según tu nuevo formato en la imagen
    hoja['B3'] = f"Registro del día: {fecha_actual}"

    # === 2. INSERCIÓN DE LOGOS POR CÓDIGO ===
    # Logo Izquierdo
    ruta_logo_izq = os.path.join(settings.BASE_DIR, 'inicio', 'plantillas', 'logo_izq.png')
    if os.path.exists(ruta_logo_izq):
        img_izq = Image(ruta_logo_izq)
        img_izq.width = 90  # Puedes ajustar el ancho en píxeles según tu diseño
        img_izq.height = 90  # Puedes ajustar el alto en píxeles
        hoja.add_image(img_izq, 'A1')  # Ubica el logo en la celda A1

    # Logo Derecho
    ruta_logo_der = os.path.join(settings.BASE_DIR, 'inicio', 'plantillas', 'logo_der.png')
    if os.path.exists(ruta_logo_der):
        img_der = Image(ruta_logo_der)
        img_der.width = 170
        img_der.height = 85
        hoja.add_image(img_der, 'F1')  # Ubica el logo en la celda F1 (o la columna que prefieras)

    # === 2. INYECCIÓN DE DATOS DE HISTORIAL ===
    historial_registros = Registro.history.all().order_by('-history_date')

    # Comenzamos a inyectar datos debajo de los encabezados (Fila 5 según tu Excel)
    fila_actual = 5
    for registro in historial_registros:
        if registro.history_type == '+':
            accion = 'Creación'
        elif registro.history_type == '~':
            accion = 'Modificación'
        else:
            accion = 'Eliminación'

        operador = registro.history_user.username if registro.history_user else 'Sistema'
        hora_mexico = timezone.localtime(registro.history_date)

        hoja.cell(row=fila_actual, column=1, value=hora_mexico.strftime('%d/%m/%y %I:%M %p'))
        hoja.cell(row=fila_actual, column=2, value=operador)
        hoja.cell(row=fila_actual, column=3, value=registro.modulo)
        hoja.cell(row=fila_actual, column=4, value=registro.tipo_registro or 'N/A')
        hoja.cell(row=fila_actual, column=5, value=accion)
        hoja.cell(row=fila_actual, column=6, value=registro.observaciones or 'Sin observaciones')

        fila_actual += 1

    # === 3. DESCARGA DEL ARCHIVO ===
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Auditoria_UAPIX_{fecha_actual.replace("/", "-")}.xlsx"'
    workbook.save(response)
    
    return response
# Hasta aqui modifico Owen

# 1. VISTA DE LOGIN (Se mantiene tu código exacto)
def login_view(request):
    if request.method == 'POST':
        username_ingresado = request.POST.get('usuario')
        password_ingresado = request.POST.get('contrasena')
        
        user = authenticate(request, username=username_ingresado, password=password_ingresado)
        
        if user is not None:
            login(request, user)
            return redirect('homePage')
        else:
            messages.error(request, "Usuario o contraseña incorrectos.")
            
    return render(request, 'inicio/login.html')

# 2. VISTA DEL HOMEPAGE (Se mantiene)
@login_required
def homePage(request):
    return render(request, 'inicio/homePage.html')

# 3. CANDADO DE SEGURIDAD (Asegura que solo el admin general cree usuarios)
def es_admin_general(user):
    return user.is_superuser

# 4. VISTA PARA AGREGAR USUARIO (Totalmente refactorizada)
@login_required
@user_passes_test(es_admin_general, login_url='homePage')
def agregar_usuario_view(request):
    if request.method == 'POST':
        # Le entregamos todos los datos que llegaron del HTML a nuestro formulario
        form = RegistroUsuarioForm(request.POST)
        
        if form.is_valid():
            # form.save() hace la magia: crea el usuario, encripta la clave y guarda los checkboxes
            form.save() 
            messages.success(request, "¡Usuario y permisos guardados exitosamente!")
            return redirect('agregar_usuario') # Recargamos la página limpia
        else:
            # Si el formulario detecta un error (ej. el nombre de usuario ya existe), avisa
            messages.error(request, "Error al crear el usuario. Verifica los datos ingresados.")
    else:
        # Si la petición es GET (solo entrar a la pantalla), mandamos el formulario vacío a Yesenia
        form = RegistroUsuarioForm()

    return render(request, 'inicio/agregar_usuario.html', {'form': form})

# 5. VISTA DE LOGOUT (Se mantiene)
def logout_view(request):
    logout(request)
    return redirect('login')
